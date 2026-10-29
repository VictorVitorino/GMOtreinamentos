"""Avaliação de Desempenho: 3 slides executivos (contexto + 2 para o Gerente Executivo), no template A&M do deck original.

Uso: python3 build_slides.py <Avaliacao_de_Desempenho.pptx> [saida.pptx]

O deck original entra só como template: dele vêm o layout "11_Texto 1 Coluna Diagonal Branco",
a barra de navegação do topo e o número de slide. Os 21 slides originais saem do arquivo final.
"""
import copy
import os
import re
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

DIST = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'dist')
LAYOUT = '11_Texto 1 Coluna Diagonal Branco'

# Paleta do tema do deck (theme2) + verde-menta e cinzas usados nos slides originais
NAVY, STEEL, SLATE, GREYBLUE, LIGHT, MINT = '002A46', '5E8AB3', '33546D', '99AAB6', 'CCD4DA', 'ACE2CC'
TXT, LINE, TINT, WHITE = '555555', 'D8D8D8', 'EEF2F5', 'FFFFFF'
BODY = 'Arial Nova Cond Light'

# Medida de texto: Liberation Sans (métrica do Arial) comprimida como a Arial Nova Cond, com folga de 4%
LIB = os.environ.get('LIB_FONTS', '/usr/share/fonts/truetype/liberation')
COND, SLACK = 0.84, 1.04
_fonts = {}
WARN = []


def _font(bold):
    k = 'Bold' if bold else 'Regular'
    if k not in _fonts:
        _fonts[k] = ImageFont.truetype(f'{LIB}/LiberationSans-{k}.ttf', 200)
    return _fonts[k]


def tw(t, size, bold=False, spc=0.0, cond=COND):
    return _font(bold).getlength(t) * size / 200 * cond * SLACK + spc * len(t)


def nlines(runs, size, width, spc=0.0, cond=COND):
    lines, cur, space = 1, 0.0, 0.0
    for t, b, _ in runs:
        for tok in re.findall(r'\S+|\s+', t):
            w = tw(tok, size, b, spc, cond)
            if tok.isspace():
                space += w
                continue
            if w > width:
                WARN.append(f'palavra não cabe em {width:.0f}pt: {tok}')
            if cur and cur + space + w > width:
                lines, cur = lines + 1, w
            else:
                cur += space + w
            space = 0.0
    return lines


def E(pt):
    return Emu(int(round(pt * 12700)))


def P(text=None, size=10, bold=False, color=TXT, lh=None, sb=0, align='l', spc=0, italic=False, runs=None, bullet=None):
    """Parágrafo; `bullet` é a cor do marcador quadrado (recuo deslocado de BULLET_IND pt)."""
    return dict(runs=runs or [(text, bold, color)], size=size, lh=lh or round(size * 1.2, 2),
                sb=sb, align=align, spc=spc, italic=italic, bullet=bullet)


BULLET_IND = 11


def paras_h(paras, w):
    return sum(nlines(p['runs'], p['size'], w - (BULLET_IND if p.get('bullet') else 0), p['spc']) * p['lh']
               + (p['sb'] if i else 0) for i, p in enumerate(paras))


def _fill_tf(tf, paras, anchor, ins=(0, 0, 0, 0)):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    bp = tf._txBody.bodyPr
    for a, v in zip(('lIns', 'tIns', 'rIns', 'bIns'), ins):
        bp.set(a, str(int(E(v))))
    bp.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[p['align']]
        para.line_spacing = Pt(p['lh'])
        para.space_before = Pt(p['sb'] if i else 0)
        para.space_after = Pt(0)
        if p.get('bullet'):
            pPr = para._p.get_or_add_pPr()
            pPr.set('marL', str(int(E(BULLET_IND))))
            pPr.set('indent', str(-int(E(BULLET_IND))))
            for tag, attrs in (('a:buClr', None), ('a:buSzPct', {'val': '60000'}),
                               ('a:buFont', {'typeface': 'Arial'}), ('a:buChar', {'char': '\u25a0'})):
                el = etree.SubElement(pPr, qn(tag), attrs or {})
                if tag == 'a:buClr':
                    etree.SubElement(el, qn('a:srgbClr'), {'val': p['bullet']})
        for t, b, c in p['runs']:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold, f.italic = BODY, Pt(p['size']), b, p['italic']
            f.color.rgb = RGBColor.from_string(c)
            rPr = r._r.get_or_add_rPr()
            rPr.set('lang', 'pt-BR')
            if p['spc']:
                rPr.set('spc', str(int(round(p['spc'] * 100))))


class Slide:
    """Slide novo no layout do template, com registro das caixas de texto para a checagem de sobreposição."""

    def __init__(self, prs, layout, ref, title, notes):
        self.sl = prs.slides.add_slide(layout)
        self.texts = []
        for ph in list(self.sl.placeholders):
            if ph.placeholder_format.idx != 19:
                ph._element.getparent().remove(ph._element)
        t = next(ph for ph in self.sl.placeholders if ph.placeholder_format.idx == 19)
        t.left, t.top, t.width, t.height = Emu(469562), E(37), Emu(11251289), E(54)
        t.text_frame.text = title
        for r in t.text_frame.paragraphs[0].runs:
            r._r.get_or_add_rPr().set('lang', 'pt-BR')
        nt = nlines([(title, False, '')], 22, 886, cond=1.0)  # título em Arial Nova Light, sem condensar
        if nt > 2:
            WARN.append(f'título com mais de 2 linhas: {title}')
        self.texts.append((f'título: {title[:30]}', 37, 37, 886, nt * 26.4, None))
        # barra de navegação do topo (aba "Avaliação de desempenho" ativa) e número do slide
        tree = self.sl.shapes._spTree
        for sh in ref.shapes:
            num = sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
            if num or sh.name in ('Espaço Reservado para Conteúdo 2', 'Retângulo 120'):
                el = copy.deepcopy(sh._element)
                for h in list(el.iter(qn('a:hlinkClick'))):
                    h.getparent().remove(h)
                tree.append(el)
        self.sl.notes_slide.notes_text_frame.text = notes

    def box(self, x, y, w, h, fill=None, line=None, lw=0.5, shadow=False, prst=MSO_SHAPE.RECTANGLE, adj=None):
        s = self.sl.shapes.add_shape(prst, E(x), E(y), E(w), E(h))
        st = s._element.find(qn('p:style'))
        if st is not None:
            s._element.remove(st)
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(lw)
        else:
            s.line.fill.background()
        if adj is not None:
            s.adjustments[0] = adj
        if shadow:
            s._element.spPr.append(etree.fromstring(
                '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '<a:outerShdw blurRad="101600" dist="25400" dir="5400000" algn="bl" rotWithShape="0">'
                '<a:srgbClr val="000000"><a:alpha val="10000"/></a:srgbClr></a:outerShdw></a:effectLst>'))
        return s

    def text(self, x, y, w, paras, h=None, anchor='t', inside=None):
        need = paras_h(paras, w)
        hh = need if h is None else h
        if need > hh + 0.5:
            WARN.append(f'texto estoura a caixa ({need:.1f} > {hh:.1f}pt): {paras[0]["runs"][0][0]}')
        tb = self.sl.shapes.add_textbox(E(x), E(y), E(w), E(hh))
        _fill_tf(tb.text_frame, paras, anchor)
        self.texts.append((paras[0]['runs'][0][0][:40], x, y, w, hh, inside))
        return need

    def shape_text(self, s, x, y, w, h, paras, ins, anchor='m'):
        need = paras_h(paras, w - ins[0] - ins[2])
        if need > h - ins[1] - ins[3] + 0.5:
            WARN.append(f'texto estoura a forma ({need:.1f}pt): {paras[0]["runs"][0][0]}')
        _fill_tf(s.text_frame, paras, anchor, ins)

    def check(self, label):
        T = self.texts
        for i, (n, x, y, w, h, box) in enumerate(T):
            if box and (x < box[0] - 0.1 or y < box[1] - 0.1 or x + w > box[0] + box[2] + 0.1 or y + h > box[1] + box[3] + 0.1):
                WARN.append(f'{label}: "{n}" sai do contêiner')
            for m, x2, y2, w2, h2, _ in T[i + 1:]:
                if x < x2 + w2 - 0.5 and x2 < x + w - 0.5 and y < y2 + h2 - 0.5 and y2 < y + h - 0.5:
                    WARN.append(f'{label}: "{n}" sobrepõe "{m}"')
        for i, c in enumerate(self.sl.shapes._spTree.iter(qn('p:cNvPr')), start=1):
            c.set('id', str(i))


X0, X1, YB = 36, 922, 486  # margens do conteúdo no layout (iguais às dos slides originais)


def cols(n, gap):
    w = (X1 - X0 - gap * (n - 1)) / n
    return [(X0 + i * (w + gap), w) for i in range(n)]


# ---------------------------------------------------------------- slide 0: contexto
# Só conteúdo do deck original: objetivo (slide 3), modelo atual do ONS (5), premissas de readequação (7)
# e ajustes essenciais (10)
CICLO_ATUAL = [
    ('Início do ciclo', 'Elaboração de metas', '2 a 4 metas ao ano. Gerente imediato e executivo aprovam.'),
    ('Meio do ciclo', 'Repactuação das metas', 'Ajustes e revisões, se necessário. Gerente imediato e executivo aprovam.'),
    ('Final do ciclo', 'Autoavaliação, avaliação final e feedback formal', 'Conduzidos pelo superior imediato.'),
]
READEQUACAO = [
    ('Modelo matricial', 'Profissionais respondem a dois vetores: o Líder de Comunidade (entrega) e o Gerente Funcional '
     '(desenvolvimento técnico e carreira), além de profissionais alocados entre diretorias.',
     'A avaliação precisa capturar ambas as dimensões.'),
    ('Papéis híbridos', 'POs, Líderes de Comunidade, Modeleiros e Desenvolvedores podem operar em squads multidisciplinares.',
     'A avaliação individual deve considerar a contribuição coletiva ao squad.'),
    ('Ciclo de inovação', 'Parte do trabalho é de exploração (Comunidades Técnicas, horizonte 2-3), com entregas de mais longo '
     'prazo e incerteza inerente incompatível com metas puramente quantitativas de curto prazo.', None),
    ('Orientação ao usuário', 'É necessário avaliar as squads sob a ótica do usuário.',
     'Entender se há pleno atendimento das expectativas tanto em qualidade, tempo de resposta e suporte.'),
    ('Diversidade de perfis', 'O DQQ inclui 14 perfis distintos (de Secretaria a Arquiteto e Especialista I), cada um com '
     'dinâmicas de trabalho e expectativas de entrega muito diferentes.', None),
]
AJUSTES = ['Adequação das competências ao contexto do modelo matricial',
           'Definição de ao menos uma meta voltada à contribuição coletiva para a comunidade',
           'Estruturação de formulário de percepção da área cliente sobre a atuação da equipe, com aplicação via Forms']


def slide_contexto(prs, layout, ref):
    S = Slide(prs, layout, ref,
              'Avaliação de Desempenho: modelo atual do ONS e premissas da readequação',
              'Proposta de modelo de avaliação de desempenho para o novo contexto organizacional. Parte do modelo atual '
              'do ONS (ciclo anual, pilares de resultados e de desenvolvimento) e das cinco premissas que exigem a '
              'readequação ao modelo matricial, resultando em três ajustes essenciais.')
    S.text(X0, 68, X1 - X0, [P('Apresenta uma proposta de modelo de avaliação de desempenho para o novo contexto organizacional, '
                              'definindo possíveis indicadores, critérios e dimensões para mensuração dos resultados das '
                              'equipes e do modelo.', 11, False, TXT, lh=13.5)])
    yp, pad = 106, 14
    # esquerda: modelo atual
    x, w = X0, 330
    card = (x, yp, w, YB - yp)
    S.box(*card, fill=WHITE, line=LINE, shadow=True)
    S.box(x, yp, w, 3, fill=GREYBLUE)
    ix, iw = x + pad, w - 2 * pad
    S.text(ix, yp + 14, iw, [P('MODELO ATUAL DO ONS', 7.5, True, STEEL, lh=9.5, spc=1.2),
                             P('Ciclo anual de avaliação', 15, True, NAVY, lh=18, sb=3)], inside=card)
    S.box(ix, yp + 52, iw, 0.75, fill=LINE)
    cy = yp + 66
    tx, wt = ix + 30, iw - 30
    pil = [('PILAR DE RESULTADOS', 'Atingimento das metas estabelecidas', 'Média de 1 a 5', NAVY),
           ('PILAR DE DESENVOLVIMENTO', 'Cesta de competências, definida de acordo com a necessidade da área',
            'Menos hábil, Hábil, Referência. Não entra na média', STEEL)]
    pw = (iw - 8) / 2
    pp = lambda k, a, b: [P(k, 7, True, MINT, lh=9, spc=0.6), P(a, 10, True, WHITE, lh=12, sb=4), P(b, 9, False, WHITE, lh=11, sb=4)]
    hp = max(paras_h(pp(k, a, b), pw - 20) for k, a, b, _ in pil) + 20
    ypil = YB - pad - hp
    etapa = lambda k, a, b: [P(k.upper(), 7, True, STEEL, lh=9, spc=0.8), P(a, 11, True, NAVY, lh=13, sb=1.5),
                             P(b, 9.5, False, TXT, lh=11.5, sb=1.5)]
    he = [paras_h(etapa(*e), wt) for e in CICLO_ATUAL]
    gape = (ypil - 26 - cy - sum(he)) / (len(he) - 1)
    for n, (k, a, b) in enumerate(CICLO_ATUAL, start=1):
        s = S.box(ix, cy + 1, 20, 20, fill=GREYBLUE)
        S.shape_text(s, ix, cy + 1, 20, 20, [P(str(n), 11, True, WHITE, lh=13, align='c')], ins=(0, 0, 0, 0))
        cy += S.text(tx, cy, wt, etapa(k, a, b), inside=card) + gape
    if gape < 10:
        WARN.append(f'contexto: etapas apertadas ({gape:.0f})')
    cy = ypil
    S.box(ix, cy - 13, iw, 0.75, fill=LINE)
    for j, (k, a, b, cor) in enumerate(pil):
        bx = ix + j * (pw + 8)
        S.box(bx, cy, pw, hp, fill=cor)
        S.text(bx + 10, cy + 10, pw - 20, pp(k, a, b), inside=card)
    # direita: premissas de readequação + ajustes
    rx = x + w + 12
    rw = X1 - rx
    S.text(rx, yp + 2, rw, [P('PREMISSAS DE READEQUAÇÃO', 7.5, True, STEEL, lh=9.5, spc=1.2),
                            P('O que a avaliação precisa capturar no novo modelo', 15, True, NAVY, lh=18, sb=3)])
    gy, gap = yp + 44, 8
    cw = (rw - gap) / 2
    ch = (YB - gy - 2 * gap) / 3
    cells = [(c, False) for c in READEQUACAO] + [(None, True)]
    for i, (c, final) in enumerate(cells):
        cx, cyy = rx + (i % 2) * (cw + gap), gy + (i // 2) * (ch + gap)
        cell = (cx, cyy, cw, ch)
        if final:
            S.box(*cell, fill=NAVY)
            paras = [P('O QUE MUDA', 7, True, MINT, lh=9, spc=1), P('Ajustes essenciais para o novo modelo', 11, True, WHITE, lh=13, sb=2)]
            for n, a in enumerate(AJUSTES, start=1):
                paras.append(P(runs=[(f'{n}  ', True, MINT), (a, False, WHITE)], size=9, lh=11, sb=3))
        else:
            t, a, b = c
            S.box(*cell, fill=WHITE, line=LINE, shadow=True)
            S.box(cx, cyy, 3, ch, fill=MINT)
            paras = [P(t, 11, True, NAVY, lh=13), P(a, 9.5, False, TXT, lh=11.5, sb=3)]
            if b:
                paras.append(P(b, 9.5, True, SLATE, lh=11.5, sb=3))
        S.text(cx + 13, cyy + 9, cw - 24, paras, inside=cell)
        if paras_h(paras, cw - 24) > ch - 16:
            WARN.append(f'contexto: célula {i + 1} estoura')
    S.check('slide contexto')


# ---------------------------------------------------------------- slide 1: roadmap
# Só conteúdo do deck original: fluxo do ciclo (slide 8), fases (11), pilar de resultados (13),
# formulário da área cliente (15) e a coluna do Gerente Executivo na matriz RACI (17)
ETAPAS = [
    ('INÍCIO DO CICLO', 'Elaboração de metas', [
        ('Metas individuais (70%)', '1 a 3 metas por ciclo, aprovadas pelo Gerente Imediato'),
        ('Meta de contribuição à Comunidade (30%)', '1 meta coletiva por ciclo, definida junto com o Líder de Comunidade'),
        ('3 metas comuns ONS', 'Definidas institucionalmente, aplicáveis a todo o ONS'),
    ], ('Consultado (C)', 'na elaboração das metas para o ciclo')),
    ('MEIO DO CICLO', 'Repactuação das metas', [
        ('Repactuação', 'Considerando mudanças de backlog, realocações entre squads e evolução dos modelos'),
        ('Ajustes', 'O Líder de Comunidade pode propor ajustes'),
    ], ('Consultado (C)', 'na repactuação das metas')),
    ('FINAL DO CICLO', 'Avaliação final', [
        ('Autoavaliação', 'Responsabilidade do avaliado'),
        ('Avaliação de desenvolvimento e de resultados', 'Pelo Gerente, com input do Líder de Comunidade (escala de 1 a 5)'),
        ('Calibração, avaliação final e feedback formal', 'Conduzidos pelo Gerente, com o Líder de Comunidade consultado'),
    ], ('Aprovador (A)', 'da calibração e avaliação final')),
    ('CICLOS SEGUINTES', 'Evolução', [
        ('Formulário de percepção da área cliente', 'Qualidade, tempo de resposta, suporte e valor percebido, em escala de 1 a 5'),
        ('NPS por Comunidade de Entrega', 'Indicador agregado de satisfação da comunidade ao longo do tempo'),
        ('Calibragem cruzada entre Comunidades', 'Fóruns entre gerentes para garantir consistência das notas no modelo matricial'),
    ], ('Informado (I)', 'no feedback do representante do negócio (formulário)')),
]


def slide_roadmap(prs, layout, ref):
    S = Slide(prs, layout, ref,
              'Roadmap de implantação: as etapas do ciclo de avaliação e o papel do Gerente Executivo em cada uma',
              'Fase 1 (início, meio e final do ciclo) já no próximo ciclo de avaliação; Fase 2 em ciclos seguintes. '
              'A faixa inferior traz a coluna do Gerente Executivo na matriz RACI.')
    C = cols(4, 8)
    # fases
    yl = 103
    f1 = (C[0][0], C[2][0] + C[2][1] - C[0][0])
    for (x, w), label, cor in ((f1, 'FASE 1  ·  JÁ NO PRÓXIMO CICLO', GREYBLUE), (C[3], 'FASE 2  ·  EVOLUÇÃO EM CICLOS SEGUINTES', MINT)):
        S.text(x, yl, w, [P(label, 8, True, NAVY, lh=10, spc=1)])
        S.box(x, yl + 13, w, 2.5, fill=cor)
    # etapas (chevrons)
    yc, hc, adj = 124, 46, 0.26
    dep = hc * adj
    for i, ((x, w), (k, nome, _, _)) in enumerate(zip(C, ETAPAS)):
        fase2 = i == 3
        s = S.box(x, yc, w, hc, fill=MINT if fase2 else NAVY, prst=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, adj=adj)
        S.shape_text(s, x, yc, w, hc, [
            P(f'{i + 1} · {k}', 7.5, True, SLATE if fase2 else MINT, lh=9.5, spc=0.8),
            P(nome, 12, True, NAVY if fase2 else WHITE, lh=14, sb=1.5),
        ], ins=(10 if i == 0 else dep + 7, 0, dep + 3, 0))
    # cartões: medir tudo antes de desenhar, para alinhar a faixa "Papel do GE"
    ycard, pad = 178, 11
    iw = C[0][1] - 2 * pad
    hd = lambda t: P(t, 11.5, True, NAVY, lh=13.5)
    ds = lambda t: P(t, 10, False, TXT, lh=12)
    corpo = [sum(paras_h([hd(a)], iw) + 2.5 + paras_h([ds(b)], iw) for a, b in it) + (len(it) - 1) * 17 for _, _, it, _ in ETAPAS]
    rod = lambda papel: [P('PAPEL DO GE (RACI)', 7, True, STEEL, lh=9, spc=1), P(papel[0], 13, True, NAVY, lh=15, sb=2),
                         P(papel[1], 9, False, TXT, lh=11, sb=1)]
    hrod = max(paras_h(rod(e[3]), iw) for e in ETAPAS) + 2 * 9
    yrod = YB - hrod
    if ycard + 12 + max(corpo) + 10 > yrod:
        WARN.append(f'roadmap: cartões sem espaço ({ycard + 12 + max(corpo) + 10:.0f} > {yrod:.0f})')
    for i, ((x, w), (_, _, itens, papel)) in enumerate(zip(C, ETAPAS)):
        card = (x, ycard, w, YB - ycard)
        S.box(*card, fill=WHITE, line=LINE, shadow=True)
        cy = ycard + 12
        for j, (a, b) in enumerate(itens):
            cy += S.text(x + pad, cy, iw, [hd(a)], inside=card) + 2.5
            cy += S.text(x + pad, cy, iw, [ds(b)], inside=card)
            if j < len(itens) - 1:
                S.box(x + pad, cy + 8, iw, 0.75, fill=LINE)
                cy += 17
        aprova = papel[0].startswith('Aprovador')
        S.box(x + 0.5, yrod, w - 1, hrod - 0.5, fill=NAVY if aprova else TINT)
        r = rod(papel)
        if aprova:
            r[0]['runs'] = [('PAPEL DO GE (RACI)', True, MINT)]
            r[1]['runs'] = [(papel[0], True, WHITE)]
            r[2]['runs'] = [(papel[1], False, WHITE)]
        S.text(x + pad, yrod + 9, iw, r, inside=card)
    S.check('slide 1')


# ---------------------------------------------------------------- slide 2: o que considerar
CADEIA = [
    ('Gerente Executivo FE', 'Diretoria (DPL)', True),
    ('Gerentes FEG e de Ferramentas (FEL, FEN, FER)', 'Gerente Executivo FE', True),
    ('Líderes de Comunidade de Entrega', 'Gerente de Ferramentas', False),
    ('Líderes de Comunidade Técnica', 'Gerente de Governança e Estratégia, com input dos Gerentes de Ferramentas', False),
    ('Profissionais DPL (PO, modeleiros, desenvolvedores)', 'Gerente de Ferramentas, com input do Líder de Comunidade de Entrega', False),
    ('Profissionais DTA alocados na FE', 'Gerente imediato da DTA, com input obrigatório do Líder de Comunidade DPL', False),
]
PREMISSAS = [  # premissas de implementação (slide 13)
    ('Capacitar a liderança', 'Para garantir o cumprimento das metas conforme as regras estabelecidas'),
    ('Formalizar e comunicar papéis', 'Gerente imediato, líder de comunidade e avaliado'),
    ('Oracle HCM como ferramenta única', 'Formalização de metas, acompanhamento e registro de feedbacks'),
    ('Alinhamento prévio', 'Entre Gerente Imediato e Líder de Comunidade antes da formalização das metas'),
    ('Governança de divergências', 'Instância de resolução entre a avaliação do Gerente Imediato e o input do Líder de Comunidade'),
    ('Monitorar a aderência', 'Nos primeiros ciclos: % de metas formalizadas no prazo e % de feedbacks registrados'),
]


def chips(S, x, y, w, items, fill, line, cor, size=8.5, h=15, gap=4, card=None):
    cx, cy = x, y
    for t in items:
        cw = tw(t, size, False) + 12
        if cx > x and cx + cw > x + w:
            cx, cy = x, cy + h + gap
        s = S.box(cx, cy, cw, h, fill=fill, line=line, lw=0.75)
        S.shape_text(s, cx, cy, cw, h, [P(t, size, False, cor, align='c')], ins=(0, 0, 0, 0))
        cx += cw + gap
    return cy + h - y


def slide_considerar(prs, layout, ref):
    S = Slide(prs, layout, ref,
              'O que o Gerente Executivo deve considerar: composição da nota, quem avalia quem e premissas de implementação',
              'Pilar de resultados e de desenvolvimento, avaliadores por papel na FE e premissas de implementação, '
              'conforme a proposta A&M.')
    C = cols(3, 10)
    yp, ym = 106, 444
    pad = 14
    painéis = [('REGRAS DO MODELO', 'Como a nota se compõe', STEEL),
               ('CADEIA DE AVALIAÇÃO', 'Quem avalia quem na FE', NAVY),
               ('PROPOSTA A&M', 'Premissas de implementação', MINT)]
    for (x, w), (k, t, cor) in zip(C, painéis):
        card = (x, yp, w, ym - yp)
        S.box(*card, fill=WHITE, line=LINE, shadow=True)
        S.box(x, yp, w, 3, fill=cor)
        S.text(x + pad, yp + 14, w - 2 * pad, [P(k, 7.5, True, STEEL, lh=9.5, spc=1.2),
                                              P(t, 15, True, NAVY, lh=18, sb=3)], inside=card)
        S.box(x + pad, yp + 52, w - 2 * pad, 0.75, fill=LINE)

    # painel 1: composição da nota
    (x, w), iw = C[0], C[0][1] - 2 * pad
    card, ix, cy = (x, yp, w, ym - yp), x + pad, yp + 64

    def pilar(cy, nome, tag):
        S.text(ix, cy, iw * 0.58, [P(nome, 11.5, True, NAVY, lh=14)], inside=card)
        S.text(ix + iw * 0.58, cy + 2.5, iw * 0.42, [P(tag, 7, True, STEEL, lh=9, spc=0.8, align='r')], inside=card)
        return cy + 20

    cy = pilar(cy, 'Pilar de Resultados', 'MÉDIA DE 1 A 5')
    w70 = iw * 0.7
    for sx, sw, f, c, t in ((ix, w70 - 1, NAVY, WHITE, '70%'), (ix + w70 + 1, iw - w70 - 1, MINT, NAVY, '30%')):
        s = S.box(sx, cy, sw, 22, fill=f)
        S.shape_text(s, sx, cy, sw, 22, [P(t, 11, True, c, lh=13, align='c')], ins=(0, 0, 0, 0))
    cy += 30
    for f, ln, a, b in ((NAVY, None, 'Metas individuais', ' · 1 a 3 por ciclo, aprovadas pelo gerente'),
                        (MINT, None, 'Contribuição à Comunidade', ' · 1 meta coletiva, com o Líder'),
                        (WHITE, GREYBLUE, '+ 3 metas comuns ONS', ' · aplicáveis a todo o ONS')):
        S.box(ix, cy + 2.5, 7, 7, fill=f, line=ln, lw=0.75)
        cy += S.text(ix + 13, cy, iw - 13, [P(runs=[(a, True, NAVY), (b, False, TXT)], size=9.5, lh=12)], inside=card) + 6
    cy += 7
    S.box(ix, cy, iw, 0.75, fill=LINE)
    cy = pilar(cy + 11, 'Pilar de Desenvolvimento', 'NÃO ENTRA NA MÉDIA')
    lw_ = 54
    for rot, its, fill, line, cor in (('Essenciais a todos', ['Confiança', 'Foco nos resultados', 'Colaboração'], LIGHT, None, NAVY),
                                       ('Gestores e líderes', ['Formação de equipes eficientes'], LIGHT, None, NAVY),
                                       ('Eletivas', ['Visão estratégica', 'Gerenciamento da complexidade', 'Orientação para a ação', 'Inovação'], WHITE, STEEL, SLATE)):
        hl = S.text(ix, cy + 1, lw_ - 4, [P(rot, 8, True, SLATE, lh=9.5)], inside=card)
        hc = chips(S, ix + lw_, cy, iw - lw_, its, fill, line, cor, card=card)
        cy += max(hl + 2, hc) + 7
    cy += 2
    S.text(ix, cy + 2, lw_ - 4, [P('Escala', 8, True, SLATE, lh=9.5)], inside=card)
    sw = (iw - lw_ - 4) / 3
    for k, (t, f, c) in enumerate((('Menos hábil', LIGHT, NAVY), ('Hábil', GREYBLUE, WHITE), ('Referência', STEEL, WHITE))):
        s = S.box(ix + lw_ + k * (sw + 2), cy, sw, 15, fill=f, prst=MSO_SHAPE.PENTAGON if k < 2 else MSO_SHAPE.RECTANGLE, adj=0.35 if k < 2 else None)
        S.shape_text(s, 0, 0, sw, 15, [P(t, 8.5, True, c, lh=10, align='c')], ins=(0, 0, 4 if k < 2 else 0, 0))
    cy += 15
    if cy > ym - 10:
        WARN.append(f'painel 1 estoura ({cy:.0f} > {ym - 10})')

    # painel 2: cadeia de avaliação
    (x, w), iw = C[1], C[1][1] - 2 * pad
    card, ix, cy = (x, yp, w, ym - yp), x + pad, yp + 62
    wa = 114
    wb = iw - wa - 10
    S.text(ix, cy, wa, [P('AVALIADO', 7, True, STEEL, lh=9, spc=1)], inside=card)
    S.text(ix + wa + 10, cy, wb, [P('AVALIADO POR', 7, True, STEEL, lh=9, spc=1)], inside=card)
    cy += 14
    alt = [max(paras_h([P(a, 9.5, True, NAVY, lh=11.5)], wa), paras_h([P(b, 9.5, False, TXT, lh=11.5)], wb)) for a, b, _ in CADEIA]
    folga = min(16, max(10, (ym - 12 - cy - sum(alt)) / len(alt)))
    for (a, b, ge), h in zip(CADEIA, alt):
        hr = h + folga
        if ge:
            S.box(x + 1, cy, w - 2, hr, fill=TINT)
        else:
            S.box(ix, cy, iw, 0.5, fill=LINE)
        S.text(ix, cy + folga / 2, wa, [P(a, 9.5, True, NAVY, lh=11.5)], inside=card)
        S.text(ix + wa + 10, cy + folga / 2, wb, [P(b, 9.5, False, SLATE if ge else TXT, lh=11.5)], inside=card)
        cy += hr
    if cy > ym - 10:
        WARN.append(f'painel 2 estoura ({cy:.0f} > {ym - 10})')

    # painel 3: premissas de implementação
    (x, w), iw = C[2], C[2][1] - 2 * pad
    card, ix, cy = (x, yp, w, ym - yp), x + pad, yp + 64
    tx, wt = ix + 28, iw - 28
    alt = [paras_h([P(a, 10.5, True, NAVY, lh=12.5)], wt) + 1.5 + paras_h([P(b, 9.5, False, TXT, lh=11.5)], wt) for a, b in PREMISSAS]
    gap = min(16, max(7, (ym - 12 - cy - sum(alt)) / (len(alt) - 1)))
    for n, (a, b) in enumerate(PREMISSAS, start=1):
        s = S.box(ix, cy + 1, 18, 18, fill=NAVY)
        S.shape_text(s, ix, cy + 1, 18, 18, [P(str(n), 10, True, WHITE, lh=12, align='c')], ins=(0, 0, 0, 0))
        ha = S.text(tx, cy, wt, [P(a, 10.5, True, NAVY, lh=12.5)], inside=card)
        hb = S.text(tx, cy + ha + 1.5, wt, [P(b, 9.5, False, TXT, lh=11.5)], inside=card)
        cy += ha + 1.5 + hb + gap
    if cy - gap > ym - 10:
        WARN.append(f'painel 3 estoura ({cy - gap:.0f} > {ym - 10})')

    # mensagem-chave
    s = S.box(X0, ym + 8, X1 - X0, YB - ym - 8, fill=NAVY)
    S.shape_text(s, X0, ym + 8, X1 - X0, YB - ym - 8, [P(runs=[
        ('Sugestão A&M:  ', True, MINT),
        ('a meta de contribuição pode ser (i) meta adicional do indivíduo, voltada à Comunidade de Entrega, ou '
         '(ii) incorporada às metas comuns ONS, com o formulário de feedback da área cliente integrando essas metas.',
         False, WHITE)], size=10.5, lh=12.5)],
        ins=(16, 0, 16, 0))
    S.check('slide 2')


def prune(prs, layout, n_orig):
    """Tira os slides originais e tudo o que só eles usavam (outros masters e layouts)."""
    lst = prs.slides._sldIdLst
    for sid in list(lst)[:n_orig]:
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    keep = layout.slide_master
    mlst = prs.part._element.find(qn('p:sldMasterIdLst'))
    for mid in list(mlst):
        rid = mid.get(qn('r:id'))
        if prs.part.related_part(rid) is not keep.part:
            prs.part.drop_rel(rid)
            mlst.remove(mid)
    for lay in list(keep.slide_layouts):
        if lay.part is not layout.part:
            keep.slide_layouts.remove(lay)


def build(src, out, slides):
    """Gera `out` com os slides de `slides` (funções prs, layout, ref) no template do deck `src`."""
    prs = Presentation(src)
    layout = next(l for m in prs.slide_masters for l in m.slide_layouts if l.name == LAYOUT)
    ref = next(s for s in prs.slides if s.slide_layout is layout and
               any(sh.name == 'Espaço Reservado para Conteúdo 2' for sh in s.shapes))
    n_orig = len(prs.slides)
    for fn in slides:
        fn(prs, layout, ref)
    prune(prs, layout, n_orig)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    for w in WARN:
        print('ATENÇÃO:', w)
    print('ok', out)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else os.path.join(DIST, 'Avaliacao_Desempenho_Roadmap_GE.pptx'),
          [slide_contexto, slide_roadmap, slide_considerar])
